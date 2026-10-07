"""术语表的导入导出：CSV 和 Excel（.xlsx）。

新格式：第一行表头写语言（英语 / 日语 / 简体中文 / 繁体中文 / 韩语，或 en / ja / zh-Hans / zh-Hant / ko），
每行是同一个词在各语言里的写法。
旧格式：两列「原文, 译文」，没有表头也行——原文按文字判断英文还是日文，译文归到简体中文。
"""
from __future__ import annotations

from app.i18n import tr

import csv
from pathlib import Path

from app.core.errors import UserFacingError
from app.core.profiles.glossary import GLOSSARY_LANGS, Row, clean_row, normalize_glossary

# 表头里认得的写法（小写比较）
_LANG_HEADERS = {
    "en": {"en", "english", "英语", "英文", "英語"},
    "ja": {"ja", "jp", "japanese", "日语", "日文", "日本語"},
    "zh-Hans": {"zh-hans", "zh", "zh-cn", "chinese", "simplified chinese", "中文", "简体中文", "简体", "簡体字中国語", "中国語"},
    "zh-Hant": {"zh-hant", "zh-tw", "traditional chinese", "繁体中文", "繁體中文", "繁体", "繁体字中国語", "繁體"},
    "ko": {"ko", "korean", "韩语", "韓語", "韩文", "韓国語", "한국어"},
}
_LEGACY_HEADER_WORDS = {"原文", "译文", "source", "target", "term", "translation", "听错", "正确", "wrong", "right", "原词", "改为"}


def _lang_of_header(cell: str) -> str | None:
    c = cell.strip().lower()
    for lang, words in _LANG_HEADERS.items():
        if c in words:
            return lang
    return None


def _read_table(path: Path) -> list[list[str]]:
    suffix = path.suffix.lower()
    rows: list[list[str]] = []
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(str(path), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            for row in ws.iter_rows(values_only=True):
                if row:
                    rows.append(["" if c is None else str(c) for c in row])
            wb.close()
        except Exception as e:  # noqa: BLE001
            raise UserFacingError(tr("读取 Excel 失败：{0}", e)) from e
    elif suffix in (".csv", ".txt", ".tsv"):
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "utf-8", "gb18030", "shift_jis"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise UserFacingError(tr("无法识别这个文件的文字编码，请在 Excel 里另存为 UTF-8 的 CSV。"))
        try:
            dialect = csv.Sniffer().sniff(text[:2000], delimiters=",\t;")
        except csv.Error:
            dialect = csv.excel
        rows = [row for row in csv.reader(text.splitlines(), dialect) if row]
    else:
        raise UserFacingError(tr("只支持 .csv 和 .xlsx 文件。"))
    return rows


def import_glossary(path: str | Path) -> list[Row]:
    """读 CSV 或 Excel，返回术语行列表。"""
    p = Path(path)
    if not p.exists():
        raise UserFacingError(tr("找不到文件：{0}", p))
    table = _read_table(p)
    if not table:
        raise UserFacingError(tr("文件里没有读到任何术语。"))
    header = table[0]
    columns = [_lang_of_header(c) for c in header]
    result: list[Row] = []
    if sum(1 for c in columns if c) >= 2:
        # 新格式：按表头的语言取列
        for row in table[1:]:
            cells = {lang: row[i] for i, lang in enumerate(columns) if lang and i < len(row)}
            r = clean_row(cells)
            if r:
                result.append(r)
    else:
        # 旧格式：原文, 译文
        pairs: dict[str, str] = {}
        for i, row in enumerate(table):
            a = row[0].strip() if len(row) > 0 else ""
            b = row[1].strip() if len(row) > 1 else ""
            if not a:
                continue
            if i == 0 and (a.lower() in _LEGACY_HEADER_WORDS or b.lower() in _LEGACY_HEADER_WORDS):
                continue
            pairs[a] = b
        result = normalize_glossary(pairs)
    if not result:
        raise UserFacingError(tr("文件里没有读到任何术语。"))
    return result


def export_glossary(path: str | Path, rows: list[Row], names: dict[str, str]) -> Path:
    """导出为 UTF-8（带 BOM）的 CSV，Excel 双击打开不乱码。表头是各语言的名字（按当前界面语言）。"""
    p = Path(path)
    with open(p, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow([names.get(lang, lang) for lang in GLOSSARY_LANGS])
        for r in rows:
            w.writerow([r.get(lang, "") for lang in GLOSSARY_LANGS])
    return p


def export_pairs_csv(path: str | Path, pairs: dict[str, str], header: tuple[str, str]) -> Path:
    """两列表（纠错规则）导出。"""
    p = Path(path)
    with open(p, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(list(header))
        for k, v in pairs.items():
            w.writerow([k, v])
    return p
